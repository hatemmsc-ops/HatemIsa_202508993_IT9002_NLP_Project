"""Apply Eman's comments and the agreed copy fixes to the October 2026 content calendar (V2).

Source (read-only): input/source/content-calendars/Content_Calendar_OCT_2026_V2_-_eMYAA_-_Eman_comments.pptx
Run from emyaa-ceo-deck/:  python3 working/apply_october_calendar_edits.py
"""
import copy
from lxml import etree as etree_
from pptx import Presentation

SRC = 'input/source/content-calendars/Content_Calendar_OCT_2026_V2_-_eMYAA_-_Eman_comments.pptx'
OUT = 'output/content-calendar/Content_Calendar_OCT_2026_V3_-_eMYAA.pptx'
FOOTER_EN = 'Source: SIFMA Capital Markets Fact Book 2026. Data as of 2025. For information only, not investment advice.'
FOOTER_AR = 'المصدر: SIFMA - Capital Markets Fact Book 2026. البيانات حتى عام 2025. لأغراض توعوية فقط، ولا تُعد نصيحة استثمارية.'

# (slide, shape id, paragraph index, expected current text, new text or None to delete)
EDITS = [
    # Investment Funds carousel, AR (slide 3) and EN (slide 4)
    (3, 40, 3, 'ما المناسب لي؟', 'أيّها يناسب أهدافي الاستثمارية؟'),
    (3, 40, 15, 'فهو يتبع نفس المؤشر', 'يتبع صندوق SPUS مؤشرًا متوافقًا مع أحكام الشريعة الإسلامية، مبنيًا على شركات مؤشر S&P 500، ويطبّق معايير فرز شرعية محددة.'),
    (3, 40, 18, 'فهم الفرق بين هيكلة', 'فهم الفرق في هيكلة الصناديق يساعدك على اتخاذ قرارات متوازنة تناسب أهدافك الاستثمارية 💡'),
    (3, 40, 19, 'يمكنك الاستكشاف', 'استكشف وتداول مختلف صناديق المؤشرات المتداولة الآن عبر تطبيق eMYAA.'),
    (4, 52, 3, 'Which one is right for me?', 'Which one fits my investment goals?'),
    (4, 52, 15, 'It tracks the same U.S. index', "SPUS tracks a Shari'ah-compliant index built from the S&P 500 universe, applying specific Shari'ah screening criteria."),
    (4, 52, 18, 'Understanding the structural', 'Knowing how funds are structured helps you make balanced decisions that fit your investment goals. 💡'),
    (4, 52, 19, 'Explore and trade various ETFs', 'Explore and trade a range of ETFs now on the eMYAA app.'),
    # U.S. Market by the Numbers carousel, AR (slide 6) and EN (slide 7)
    (6, 79, 1, 'دليلك المختصر', 'دليلك المختصر لسوق الأسهم الأمريكية بالأرقام 📊'),
    (6, 79, 4, 'المركز الأول عالمياً', 'أكبر سوق أسهم في العالم 🌐'),
    (6, 79, 5, 'استحوذت أسواق الأسهم', 'شكّلت أسواق الأسهم الأمريكية 43.7% من إجمالي القيمة السوقية للأسهم عالميًا في عام 2025.'),
    (6, 79, 9, 'بلغت القيمة السوقية', 'بلغت القيمة السوقية لأسواق الأسهم الأمريكية 68.9 تريليون دولار في العام الماضي، أي ما يعادل 4.4 أضعاف حجم السوق الصينية، ثاني أكبر سوق أسهم في العالم.'),
    (6, 79, 10, '(ما يعادل 4.4', None),
    (6, 79, 14, 'أكثر من 5,500', 'أكثر من 5,500 شركة مدرجة في بورصتي نيويورك (NYSE) وناسداك (NASDAQ)، ما يمنحك خيارات واسعة بين القطاعات والشركات.'),
    (7, 90, 1, 'Your quick guide', 'Your quick guide to the U.S. stock market, by the numbers 📊'),
    (7, 90, 4, 'Unrivaled Global #1', "The World's Largest Equity Market 🌐"),
    (7, 90, 5, 'U.S. stock markets account', 'U.S. equity markets represented 43.7% of global equity market capitalization in 2025.'),
    (7, 90, 9, 'The market capitalization', "The market capitalization of U.S. equity markets reached $68.9 trillion last year, 4.4 times the size of China's market, the world's second largest."),
    (7, 90, 10, '(equivalent to 4.4', None),
    (7, 90, 14, 'More than 5,500', 'More than 5,500 companies are listed on the New York Stock Exchange (NYSE) and NASDAQ, giving you a wide choice of sectors and companies.'),
    # What Happens After You Tap Buy, video (slide 9)
    (9, 115, 2, 'ماذا يحدث بعد', 'ماذا يحدث بعد ما تضغط «Buy»؟'),
    (9, 115, 8, 'من قسم Trade', 'من قسم Trade، تقدر تتابع حركة السعر وتشوف الرسم البياني (Chart) للسهم.'),
    (9, 115, 10, 'من لحظة الضغط', 'من لحظة الضغط على «Buy»… تقدر تتابع كل استثمارك من خلال تطبيق eMYAA.'),
    (9, 117, 3, 'Head to **Orders**', 'Head to Orders at the bottom of the app to check your order status and see whether the market is open or when it’s scheduled to open.'),
    (9, 117, 5, 'Go to Trade', 'Go to Trade to track the stock’s price movement and view its chart.'),
]
# footer for every slide of the carousel, added after Post 5:
# (slide, shape id, label paragraph to copy, last paragraph to follow, label, text)
FOOTERS = [(6, 79, 16, 18, 'الفوتر (على جميع الشرائح):', FOOTER_AR),
           (7, 90, 16, 18, 'Footer (on every slide):', FOOTER_EN)]


def set_para(p, text):
    runs = p.runs
    runs[0].text = text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)


prs = Presentation(SRC)
slides = list(prs.slides)
shape = lambda n, sid: next(sh for sh in slides[n - 1].shapes if sh.shape_id == sid)

# keep paragraph objects by their original index before any insert or delete
paras = {}
for n, sid, i, *_ in EDITS:
    paras[(n, sid, i)] = shape(n, sid).text_frame.paragraphs[i]

# footer anchors, also taken before the edits shift paragraph positions
foot = [(shape(n, sid).text_frame, shape(n, sid).text_frame.paragraphs[li - 1], shape(n, sid).text_frame.paragraphs[li],
         shape(n, sid).text_frame.paragraphs[la], label, text) for n, sid, li, la, label, text in FOOTERS]

for n, sid, i, expect, new in EDITS:
    p = paras[(n, sid, i)]
    assert p.text.strip().startswith(expect), (n, sid, i, p.text)
    if new is None:
        p._p.getparent().remove(p._p)
    else:
        set_para(p, new)

from pptx.text.text import _Paragraph
for tf, gap_p, label_p, last_p, label, text in foot:
    gap = copy.deepcopy(gap_p._p)                             # empty line between posts
    lab, txt = copy.deepcopy(label_p._p), copy.deepcopy(last_p._p)
    last_p._p.addnext(gap); gap.addnext(lab); lab.addnext(txt)
    set_para(_Paragraph(lab, tf), label)
    t = _Paragraph(txt, tf); set_para(t, text); t.runs[0].font.italic = True

# the footer adds three lines: 13 pt to 11.5 pt keeps both carousels inside their text box
from pptx.util import Pt
from pptx.enum.text import MSO_ANCHOR
for n, sid, *_ in FOOTERS:
    tf = shape(n, sid).text_frame
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for para in tf.paragraphs:
        for r in para.runs:
            r.font.size = Pt(11.5)

# remove every review comment (Eman's) so the file can go to the agency
from pptx.oxml.ns import qn as _qn
for sl in slides:
    for rid, rel in list(sl.part.rels.items()):
        if rel.reltype.endswith('/comments'):
            sl.part.drop_rel(rid)
    for ext in list(sl._element.iter(_qn('p:ext'))):
        if any(etree_.QName(c).localname == 'commentRel' for c in ext):
            lst = ext.getparent(); lst.remove(ext)
            if len(lst) == 0:
                lst.getparent().remove(lst)
for rid, rel in list(prs.part.rels.items()):
    if rel.reltype.endswith('/authors') or rel.reltype.endswith('/commentAuthors'):
        prs.part.drop_rel(rid)
prs.core_properties.last_modified_by = 'Hatem Isa Hatem'

import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
prs.save(OUT)
print('saved', OUT, len(EDITS), 'edits,', len(FOOTERS), 'footers')
