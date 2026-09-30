"""Make the CEO Progress Update (the Canva export of 01.10.2026) consistent and Canva-safe.

1. Marketing slides with the navy header band get the same header as the rest of the deck
   (blue bar, blue title, small caps subtitle, Ajyad logo top right, eMYAA logo bottom right).
2. Every SVG logo or icon is replaced with a high-resolution PNG (Canva shows SVG image fills as a blue box).
3. Every gradient fill on a shape is replaced with a PNG of the same gradient (Canva turns them light blue);
   gradient text becomes solid blue.
Source (read-only): input/source/eMYAA_Progress_Update_CEO_-_01.10.2026.pptx
Run from emyaa-ceo-deck/:  python3 working/canva_fix_progress_update.py
"""
import copy, io, math, re
import cairosvg
from PIL import Image
from lxml import etree
from pptx import Presentation
from pptx.oxml.ns import qn

SRC = 'input/source/eMYAA_Progress_Update_CEO_-_01.10.2026.pptx'
OUT = 'output/ceo-update/eMYAA_Progress_Update_CEO_-_01.10.2026_-_Canva-ready.pptx'
EMU = 914400
TITLE_BLUE = '1F3B8C'
SVG_NS = 'http://schemas.microsoft.com/office/drawing/2016/SVG/main'

prs = Presentation(SRC)
slides = list(prs.slides)
REF = slides[6]                     # "Next App" slide: the deck's standard header
REF_IDS = {'ajyad': 2, 'title_a': 4, 'title': 6, 'sub': 9, 'emyaa': 54, 'bar': 55}


def shape_by_id(slide, sid):
    return next(sh for sh in slide.shapes if sh.shape_id == sid)


def copy_into(src_slide, dst_slide, el):
    """Deep-copy an element into another slide, carrying its image relationships."""
    el = copy.deepcopy(el)
    for node in el.iter():
        for att in (qn('r:embed'), qn('r:link'), qn('r:id')):
            rid = node.get(att)
            if rid and rid in src_slide.part.rels:
                rel = src_slide.part.rels[rid]
                node.set(att, dst_slide.part.relate_to(rel._target, rel.reltype))
    dst_slide.shapes._spTree.append(el)
    return el


def renumber_ids(slide):
    for i, e in enumerate(slide._element.iter(qn('p:cNvPr')), 2):
        e.set('id', str(i))


def set_group_text(el, text):
    runs = el.findall('.//' + qn('a:r'))
    runs[0].find(qn('a:t')).text = text
    for r in runs[1:]:
        r.getparent().remove(r)


# ---------- 1. one header style ----------
BAND = {  # slide number: (title, subtitle)
    14: ('Marketing Efforts', None),
    15: ('Marketing Results', None),
    16: ('English vs Arabic', None),
    17: ('Abdulelah Masterclass & Google Ads', 'INFLUENCER VIDEO AND CAMPAIGN UPDATE  |  FOR WEALTHTECH ONLY'),
    18: ('Google Ads SEO Ranking', 'SCREENSHOT OF THE GOOGLE ADS RANKING'),
    24: ('Google Ads and Installs', None),
}
for n, (title, sub) in BAND.items():
    s = slides[n - 1]
    tree = s.shapes._spTree
    headline = None
    for sh in list(s.shapes):
        top, bottom = sh.top / EMU, (sh.top + sh.height) / EMU
        is_footer_logo = abs(sh.left / EMU - 17.4) < 0.1 and top > 10.2
        if bottom <= 2.25 or is_footer_logo:
            if sh.shape_type == 6:
                txt = ' '.join(x.text_frame.text for x in sh.shapes if x.has_text_frame and x.text_frame.text.strip())
                if len(txt) > 20:
                    headline = txt.strip()
            tree.remove(sh._element)
    # move content up to where the standard header ends (the Abdulelah video is resized to fit)
    rest = [sh for sh in s.shapes]
    if rest:
        media = [sh for sh in rest if sh.shape_type == 16]
        for m in media:
            if m.top / EMU < 2.35:
                h = 10.15 * EMU - 2.35 * EMU
                m.width = int(m.width * h / m.height); m.height = int(h); m.top = int(2.35 * EMU)
        tops = [sh.top for sh in rest if sh.shape_type != 16 and sh.top / EMU < 10.2]
        if tops:
            delta = min(tops) - int(2.4 * EMU)
            if delta > 0.2 * EMU:
                for sh in rest:
                    if sh.shape_type != 16:
                        sh.top = sh.top - delta
    for key in ('ajyad', 'title_a', 'title', 'sub', 'emyaa', 'bar'):
        el = copy_into(REF, s, shape_by_id(REF, REF_IDS[key])._element)
        if key == 'title':
            set_group_text(el, title)
        if key == 'sub':
            set_group_text(el, (sub or headline or '').upper())
    renumber_ids(s)


# ---------- 2. SVG images to PNG ----------
svg_cache = {}
def svg_png(part, px):
    key = (part.partname, px)
    if key not in svg_cache:
        png = cairosvg.svg2png(bytestring=part.blob, output_width=px)
        svg_cache[key] = png
    return svg_cache[key]


def ext_px(blip):
    """Display width of the shape that holds this blip, in pixels at ~300 dpi."""
    sp = blip
    while sp is not None and etree.QName(sp).localname not in ('sp', 'pic'):
        sp = sp.getparent()
    ext = sp.find('.//' + qn('a:ext'))
    w = int(ext.get('cx')) / EMU if ext is not None else 2
    return max(600, min(3000, int(w * 300)))


# ---------- 3. gradients to PNG ----------
def hexc(h):
    return tuple(int(h[i:i + 2], 16) for i in (0, 2, 4))


def grad_png(grad, w_emu, h_emu):
    stops = []
    for gs in grad.iter(qn('a:gs')):
        c = gs.find(qn('a:srgbClr'))
        stops.append((int(gs.get('pos')) / 100000, hexc(c.get('val'))))
    stops.sort()
    W = 600
    H = max(8, min(2400, int(W * h_emu / max(w_emu, 1))))
    if H > 1200:
        W = int(W * 1200 / H); H = 1200
    lin, path = grad.find(qn('a:lin')), grad.find(qn('a:path'))
    if lin is not None:
        ang = math.radians(int(lin.get('ang', '0')) / 60000)
        dx, dy = math.cos(ang), math.sin(ang)
        span = abs(dx) * W + abs(dy) * H
        x0 = 0 if dx >= 0 else W; y0 = 0 if dy >= 0 else H
        f = lambda x, y: ((x - x0) * dx + (y - y0) * dy) / span
    else:
        ftr = path.find(qn('a:fillToRect'))
        cx = int(ftr.get('l', 50000)) / 100000 if ftr is not None else 0.5
        cy = int(ftr.get('t', 50000)) / 100000 if ftr is not None else 0.5
        far = max(math.hypot(ax - cx, ay - cy) for ax in (0, 1) for ay in (0, 1))
        f = lambda x, y: math.hypot(x / W - cx, y / H - cy) / far
    img = Image.new('RGB', (W, H))
    px = img.load()
    for y in range(H):
        for x in range(W):
            t = min(1, max(0, f(x, y)))
            for i in range(len(stops) - 1):
                (p0, c0), (p1, c1) = stops[i], stops[i + 1]
                if t <= p1 or i == len(stops) - 2:
                    u = 0 if p1 == p0 else min(1, max(0, (t - p0) / (p1 - p0)))
                    px[x, y] = tuple(int(c0[k] + (c1[k] - c0[k]) * u) for k in range(3))
                    break
    buf = io.BytesIO(); img.save(buf, 'PNG'); buf.seek(0)
    return buf


counts = {'svg': 0, 'grad_shape': 0, 'grad_text': 0}
# the deck's card gradient (radial, light blue corner to navy): used for every card so they all match
STD = next(g for g in slides[1]._element.iter(qn('a:gradFill'))
           if etree.QName(g.getparent()).localname == 'spPr' and g.find(qn('a:path')) is not None)

for s in slides:
    part = s.part
    # SVG blips
    for blip in list(s._element.iter(qn('a:blip'))):
        svgb = blip.find('.//{%s}svgBlip' % SVG_NS)
        if svgb is None:
            continue
        svg_part = part.related_part(svgb.get(qn('r:embed')))
        png = svg_png(svg_part, ext_px(blip))
        _, rid = part.get_or_add_image_part(io.BytesIO(png))
        blip.set(qn('r:embed'), rid)
        ext_lst = blip.find(qn('a:extLst'))
        for ext in list(ext_lst):
            if ext.find('{%s}svgBlip' % SVG_NS) is not None:
                ext_lst.remove(ext)
        if len(ext_lst) == 0:
            blip.remove(ext_lst)
        counts['svg'] += 1
    # gradients
    for grad in list(s._element.iter(qn('a:gradFill'))):
        par = grad.getparent()
        if etree.QName(par).localname == 'rPr':
            sf = etree.Element(qn('a:solidFill'))
            etree.SubElement(sf, qn('a:srgbClr'), val=TITLE_BLUE)
            par.replace(grad, sf)
            counts['grad_text'] += 1
        elif etree.QName(par).localname == 'spPr':
            ext = par.find(qn('a:xfrm') + '/' + qn('a:ext'))
            if b'2B4AA0' in etree.tostring(grad):      # Client Experience cards: match the other cards
                src = STD
            else:
                src = grad
            _, rid = part.get_or_add_image_part(grad_png(src, int(ext.get('cx')), int(ext.get('cy'))))
            bf = etree.Element(qn('a:blipFill'))
            etree.SubElement(bf, qn('a:blip')).set(qn('r:embed'), rid)
            st = etree.SubElement(bf, qn('a:stretch')); etree.SubElement(st, qn('a:fillRect'))
            par.replace(grad, bf)
            counts['grad_shape'] += 1

prs.save(OUT)
print('saved', OUT, counts)
