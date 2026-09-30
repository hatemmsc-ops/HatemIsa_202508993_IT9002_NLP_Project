"""Rebuild the eMYAA Daily Performance Dashboard PDFs as an editable PowerPoint.

Every page becomes one slide: text as real text boxes (same position, size, colour),
cards / bars / dividers as native rectangles, icons as high-resolution crops, the header
as its original background image, and the eMYAA logo from the brand assets.
Run from emyaa-ceo-deck/:  python3 working/dashboard_to_pptx.py
"""
import io, os, re
import pymupdf
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE

SOURCES = [  # newest first, matching the PDF order
    'input/analytics/dashboards/eMYAA_-_Daily_Performance_Dashboard_34.pdf',   # 30 Sep 2026
    'input/analytics/dashboards/eMYAA_-_Daily_Performance_Dashboard-1.pdf',    # 29 Sep back to 20 May 2026
]
OUT = 'output/dashboard-pptx/eMYAA_-_Daily_Performance_Dashboard.pptx'
WORK = 'working/dashboard-assets'
PAGE_W = 1440.0
SLIDE_W = Emu(12192000)                 # 13.333 in
EMU_PER_PT = 12192000 / PAGE_W          # PDF point -> EMU
FONT_SCALE = 960.0 / PAGE_W             # PDF font pt -> slide font pt
HEADER_H = 90.0
FONT = 'Exo 2'
FONT_FILES = {False: 'input/brand-assets/fonts/Exo2.ttf', True: 'input/brand-assets/fonts/Exo2Bold.ttf'}

E = lambda v: Emu(int(round(v * EMU_PER_PT)))
hexrgb = lambda c: RGBColor.from_string('%02X%02X%02X' % tuple(round(x * 255) for x in c))


def build_assets(page):
    """Header background and icon crops, taken once from the newest page (layout is identical on all pages)."""
    os.makedirs(WORK, exist_ok=True)
    doc = page.parent
    # header: the full-page background image, cropped to the visible header band
    info = max(page.get_image_info(xrefs=True), key=lambda i: i['width'] * i['height'])
    img = Image.open(io.BytesIO(doc.extract_image(info['xref'])['image'])).convert('RGB')
    x0, y0, x1, y1 = info['bbox']
    sx, sy = img.width / (x1 - x0), img.height / (y1 - y0)
    img.crop((round((0 - x0) * sx), round((0 - y0) * sy), round((PAGE_W - x0) * sx), round((HEADER_H - y0) * sy))).save(f'{WORK}/header.png')
    # icons: the tinted icon squares, rendered at high resolution
    icons = []
    for dr in page.get_drawings():
        r = dr['rect']
        if [i[0] for i in dr['items']] == ['re'] and 28 <= r.width <= 31.5 and 28 <= r.height <= 31.5 and (dr.get('fill_opacity') or 0) > 0:
            clip = pymupdf.Rect(r.x0 - 1, r.y0 - 1, r.x1 + 1, r.y1 + 1)
            path = f'{WORK}/icon_{len(icons):02d}.png'
            page.get_pixmap(clip=clip, dpi=432).save(path)
            icons.append((clip, path))
    return f'{WORK}/header.png', icons


def spaced(text):
    """'D O W N L O A D S  &  R A T I N G S' -> 'DOWNLOADS & RATINGS' (letter-spaced headings)."""
    t = text.strip()
    if re.fullmatch(r'(\S( {1,2}|$))+', t) and len(t) > 6 and ' ' in t and all(len(w) == 1 for w in t.split()):
        return re.sub(r' (?! )', '', t.replace('  ', '\x00')).replace('\x00', ' ')
    return None


def add_text(slide, span):
    x0, y0, x1, y1 = span['bbox']
    bold = 'Bold' in span['font']
    size = span['size'] * FONT_SCALE
    text = span['text']
    spc = None
    collapsed = spaced(text)
    if collapsed:
        f = ImageFont.truetype(FONT_FILES[bold], 200)
        natural = f.getlength(collapsed) / 200 * span['size']
        gap = max((x1 - x0) - natural, 0) / max(len(collapsed) - 1, 1)
        spc = int(round(gap * FONT_SCALE * 100))
        text = collapsed
    if 'Type3' in span['font'] and text.strip() == '':
        return
    right = x1 > 1300 and y1 < HEADER_H
    width = (x1 - x0) * 1.25 + 10
    left = x1 - width if right else x0
    tb = slide.shapes.add_textbox(E(left), E(y0), E(width), E(y1 - y0))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
    run = p.add_run()
    run.text = text
    run.font.name = FONT
    run.font.size = Pt(round(size * 2) / 2)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string('%06X' % span['color'])
    if spc:
        run.font._element.set('spc', str(spc))


def add_page(prs, page, header_png, icons, logo_png):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = RGBColor(0xF8, 0xFA, 0xFC)
    slide.shapes.add_picture(header_png, 0, 0, SLIDE_W, E(HEADER_H))
    for dr in page.get_drawings():
        r = dr['rect']
        if [i[0] for i in dr['items']] != ['re'] or not dr.get('fill') or (dr.get('fill_opacity') or 0) == 0:
            continue
        if r.width >= PAGE_W - 1 or r.width * r.height < 30:
            continue
        if 28 <= r.width <= 31.5 and 28 <= r.height <= 31.5:
            continue  # icon squares: drawn by the icon crops
        shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(r.x0), E(r.y0), E(r.width), E(r.height))
        shp.fill.solid()
        shp.fill.fore_color.rgb = hexrgb(dr['fill'])
        shp.line.fill.background()
        shp.shadow.inherit = False
    slide.shapes.add_picture(logo_png, E(55.1), E(16.6), E(142.0), E(50.5))
    for clip, path in icons:
        slide.shapes.add_picture(path, E(clip.x0), E(clip.y0), E(clip.width), E(clip.height))
    for b in page.get_text('dict')['blocks']:
        for line in b.get('lines', []):
            for span in line['spans']:
                if span['text'].strip():
                    add_text(slide, span)
    date = re.search(r'Report Date:\s*([^\n]+)', page.get_text())
    slide.notes_slide.notes_text_frame.text = 'Rebuilt from the eMYAA Daily Performance Dashboard PDF, report date ' + (date.group(1).strip() if date else 'unknown') + '.'
    return date.group(1).strip() if date else '?'


def main():
    newest = pymupdf.open(SOURCES[0])[0]
    header_png, icons = build_assets(newest)
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_W, Emu(6858000)
    logo = 'working/charts/emyaa_white.png'
    dates = []
    for src in SOURCES:
        for page in pymupdf.open(src):
            dates.append(add_page(prs, page, header_png, icons, logo))
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    prs.save(OUT)
    print(len(dates), 'slides:', dates[0], '...', dates[-1], '| icons', len(icons))


if __name__ == '__main__':
    main()
