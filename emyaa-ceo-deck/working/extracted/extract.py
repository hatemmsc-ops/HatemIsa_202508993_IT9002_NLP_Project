import json, re, zipfile, collections
from pptx import Presentation
from pptx.util import Emu
SRC='input/source/eMYAA-Marketing-Update-01.10.2026-Enhanced.pptx'
p=Presentation(SRC)
inv={'file':SRC,'slide_width_in':p.slide_width/914400,'slide_height_in':p.slide_height/914400,'slides':[]}
for i,s in enumerate(p.slides,1):
    sl={'n':i,'layout':s.slide_layout.name,'shapes':[]}
    for sh in s.shapes:
        d={'name':sh.name,'type':str(sh.shape_type),'x':round(sh.left/914400,2) if sh.left is not None else None,'y':round(sh.top/914400,2) if sh.top is not None else None,'w':round(sh.width/914400,2) if sh.width else None,'h':round(sh.height/914400,2) if sh.height else None}
        if sh.has_text_frame and sh.text_frame.text.strip(): d['text']=sh.text_frame.text
        if getattr(sh,'has_table',False) and sh.has_table: d['table']=[[c.text for c in r.cells] for r in sh.table.rows]
        if getattr(sh,'has_chart',False) and sh.has_chart:
            c=sh.chart; d['chart']={'type':str(c.chart_type),'title':c.chart_title.text_frame.text if c.has_title else None,
              'categories':list(c.plots[0].categories),'series':[{'name':se.name,'values':list(se.values)} for pl in c.plots for se in pl.series]}
        if sh.shape_type==13: d['image']=sh.image.filename if hasattr(sh,'image') else 'picture'
        sl['shapes'].append(d)
    inv['slides'].append(sl)
z=zipfile.ZipFile(SRC)
xml=''.join(z.read(n).decode('utf8','ignore') for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
inv['colors_srgb']=collections.Counter(re.findall(r'srgbClr val="([0-9A-Fa-f]{6})"',xml)).most_common(15)
inv['fonts']=collections.Counter(re.findall(r'typeface="([^"]+)"',xml)).most_common(10)
inv['media']=[(n,z.getinfo(n).file_size) for n in z.namelist() if n.startswith('ppt/media/')]
inv['notes_slides']=[n for n in z.namelist() if 'notesSlide' in n]
json.dump(inv,open('working/extracted/inventory.json','w'),indent=1,ensure_ascii=False)
print(json.dumps({k:inv[k] for k in ['slide_width_in','slide_height_in','colors_srgb','fonts','media','notes_slides']},indent=0))
